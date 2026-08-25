import math
from io import BytesIO
from sqlalchemy.orm import Session
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

from app.models.site import Site
from app.models.site_scores import SiteScore
from app.models.environmental import SiteEnvironmentalData
from app.models.solar_prediction import SolarPrediction
from app.models.wind_prediction import WindPrediction
from app.models.energy_forecast import EnergyForecast
from geoalchemy2.shape import to_shape

# Disclosures
from app.services.suitability_engine import (
    SUITABILITY_FORMULA_USED,
    SUITABILITY_FORMULA_SOURCE,
    ENVIRONMENTAL_SCORE_NOTE,
    INFRASTRUCTURE_SCORE_NOTE,
    ECONOMIC_SCORE_NOTE
)
from app.api.predictions import (
    WIND_CAP_FACTOR_NOTE,
    SOLAR_CAP_FACTOR_NOTE,
    WIND_TURB_FORMULA_USED,
    WIND_TURB_FORMULA_SOURCE
)
from app.services.forecasting_engine import (
    FORECAST_FORMULA_USED,
    FORECAST_FORMULA_SOURCE,
    GRID_CONTRIBUTION_NOTE,
    ELECTRICITY_RATE_SOURCE,
    DEGRADATION_RATE_SOURCE
)

def fetch_report_data(site_id: int, db: Session):
    site = db.query(Site).filter(Site.id == site_id).first()
    if not site:
        raise ValueError("Site not found")
        
    geom = to_shape(site.geom)
    lat, lon = geom.y, geom.x
    
    score = db.query(SiteScore).filter(SiteScore.site_id == site_id).first()
    env = db.query(SiteEnvironmentalData).filter(SiteEnvironmentalData.site_id == site_id).first()
    solar_pred = db.query(SolarPrediction).filter(SolarPrediction.site_id == site_id).order_by(SolarPrediction.created_at.desc()).first()
    wind_pred = db.query(WindPrediction).filter(WindPrediction.site_id == site_id).order_by(WindPrediction.created_at.desc()).first()
    forecast = db.query(EnergyForecast).filter(EnergyForecast.site_id == site_id).first()
    
    return {
        "site": site,
        "lat": lat,
        "lon": lon,
        "score": score,
        "env": env,
        "solar_pred": solar_pred,
        "wind_pred": wind_pred,
        "forecast": forecast
    }

def generate_pdf_report(site_id: int, report_type: str, db: Session) -> bytes:
    data = fetch_report_data(site_id, db)
    site = data["site"]
    
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=54, leftMargin=54, topMargin=54, bottomMargin=54)
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#1A365D'),
        spaceAfter=15
    )
    section_style = ParagraphStyle(
        'DocSection',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=14,
        leading=18,
        textColor=colors.HexColor('#2B6CB0'),
        spaceBefore=12,
        spaceAfter=6
    )
    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#2D3748')
    )
    disclosure_style = ParagraphStyle(
        'DocDisclosure',
        parent=styles['Italic'],
        fontName='Helvetica-Oblique',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#4A5568')
    )
    
    elements = []
    
    # Header block
    elements.append(Paragraph(f"Site Report: {site.name} (ID: {site_id})", title_style))
    elements.append(Paragraph(f"Coordinates: {data['lat']:.5f}, {data['lon']:.5f}", body_style))
    elements.append(Paragraph(f"Report Type: {report_type.replace('_', ' ').title()}", body_style))
    elements.append(Spacer(1, 15))
    
    if report_type == "site_assessment":
        elements.append(Paragraph("1. Deployment Suitability Breakdown", section_style))
        score = data["score"]
        score_val = score.overall_deployment_score if score else 0.0
        cat_val = score.suitability_category if score else "N/A"
        
        table_data = [
            [Paragraph("<b>Metric</b>", body_style), Paragraph("<b>Score / Value</b>", body_style)],
            [Paragraph("Overall Score", body_style), Paragraph(str(score_val), body_style)],
            [Paragraph("Suitability Category", body_style), Paragraph(cat_val, body_style)],
            [Paragraph("Solar Sub-score", body_style), Paragraph(str(score.solar_score if score else "N/A"), body_style)],
            [Paragraph("Wind Sub-score", body_style), Paragraph(str(score.wind_score if score else "N/A"), body_style)],
            [Paragraph("Geographic Score", body_style), Paragraph(str(score.geographic_score if score else "N/A"), body_style)],
            [Paragraph("Infrastructure Score", body_style), Paragraph(str(score.infrastructure_score if score else "N/A"), body_style)],
            [Paragraph("Environmental Score", body_style), Paragraph(str(score.environmental_score if score else "N/A"), body_style)],
            [Paragraph("Economic Score", body_style), Paragraph(str(score.economic_score if score else "N/A"), body_style)]
        ]
        
        t = Table(table_data, colWidths=[200, 200])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#EDF2F7')),
            ('ALIGN', (0,0), (-1,-1), 'LEFT'),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E0')),
            ('PADDING', (0,0), (-1,-1), 6),
        ]))
        elements.append(t)
        elements.append(Spacer(1, 15))
        
        elements.append(Paragraph("2. Data Sources & Regulatory Disclosures", section_style))
        elements.append(Paragraph(f"<b>Formula Used</b>: {SUITABILITY_FORMULA_USED} ({SUITABILITY_FORMULA_SOURCE})", body_style))
        elements.append(Spacer(1, 6))
        elements.append(Paragraph(f"<b>Economic score proxy disclosure</b>: {ECONOMIC_SCORE_NOTE}", disclosure_style))
        elements.append(Spacer(1, 6))
        elements.append(Paragraph(f"<b>Environmental score Copernicus blocker note</b>: {ENVIRONMENTAL_SCORE_NOTE}", disclosure_style))
        elements.append(Spacer(1, 6))
        elements.append(Paragraph(f"<b>Infrastructure proximity proxy note</b>: {INFRASTRUCTURE_SCORE_NOTE}", disclosure_style))
        
    elif report_type == "potential":
        elements.append(Paragraph("1. Solar and Wind Generation Potential", section_style))
        s_pred = data["solar_pred"]
        w_pred = data["wind_pred"]
        
        table_data = [
            [Paragraph("<b>Resource Indicator</b>", body_style), Paragraph("<b>Value</b>", body_style)],
            [Paragraph("Annual Solar Incident Irradiance", body_style), Paragraph(f"{s_pred.annual_irradiance if s_pred else 'N/A'} kWh/m²", body_style)],
            [Paragraph("Expected Annual Solar Output (1kW reference)", body_style), Paragraph(f"{s_pred.expected_energy_output if s_pred else 'N/A'} kWh/year", body_style)],
            [Paragraph("Solar Capacity Factor", body_style), Paragraph(str(s_pred.capacity_factor if s_pred else "N/A"), body_style)],
            [Paragraph("Average Wind Speed", body_style), Paragraph(f"{w_pred.average_wind_speed if w_pred else 'N/A'} m/s", body_style)],
            [Paragraph("Wind Power Density (WPD)", body_style), Paragraph(f"{w_pred.wind_power_density if w_pred else 'N/A'} W/m²", body_style)],
            [Paragraph("Wind Turbulence Intensity (TI)", body_style), Paragraph(str(w_pred.turbulence_intensity if w_pred else "N/A"), body_style)],
            [Paragraph("Wind Capacity Factor (2MW reference)", body_style), Paragraph(str(w_pred.capacity_factor if w_pred else "N/A"), body_style)],
            [Paragraph("Expected Annual Wind Energy (2MW reference)", body_style), Paragraph(f"{w_pred.expected_annual_energy_production if w_pred else 'N/A'} kWh/year", body_style)]
        ]
        
        t = Table(table_data, colWidths=[220, 180])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#EDF2F7')),
            ('ALIGN', (0,0), (-1,-1), 'LEFT'),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E0')),
            ('PADDING', (0,0), (-1,-1), 6),
        ]))
        elements.append(t)
        elements.append(Spacer(1, 15))
        
        elements.append(Paragraph("2. Performance Disclosures & Physics Constants", section_style))
        elements.append(Paragraph(f"<b>Performance Formula</b>: {FORECAST_FORMULA_USED} ({FORECAST_FORMULA_SOURCE})", body_style))
        elements.append(Spacer(1, 6))
        elements.append(Paragraph(f"<b>Solar CF proxy note</b>: {SOLAR_CAP_FACTOR_NOTE}", disclosure_style))
        elements.append(Spacer(1, 6))
        elements.append(Paragraph(f"<b>Wind CF proxy note</b>: {WIND_CAP_FACTOR_NOTE}", disclosure_style))
        elements.append(Spacer(1, 6))
        elements.append(Paragraph(f"<b>Wind turbulence assumption</b>: {WIND_TURB_FORMULA_USED} ({WIND_TURB_FORMULA_SOURCE})", disclosure_style))

    elif report_type == "feasibility":
        elements.append(Paragraph("1. Economic Feasibility & Grid Contribution", section_style))
        f = data["forecast"]
        
        table_data = [
            [Paragraph("<b>Economic Metric</b>", body_style), Paragraph("<b>Value</b>", body_style)],
            [Paragraph("Solar Projected Year-1 Revenue", body_style), Paragraph(f"${f.solar_annual_revenue if f else 'N/A'}", body_style)],
            [Paragraph("Wind Projected Year-1 Revenue", body_style), Paragraph(f"${f.wind_annual_revenue if f else 'N/A'}", body_style)],
            [Paragraph("Combined Projected Year-1 Revenue", body_style), Paragraph(f"${f.combined_annual_revenue if f else 'N/A'}", body_style)],
            [Paragraph("Electricity Price Utility Rate", body_style), Paragraph(f"${f.electricity_rate_usd_kwh if f else 'N/A'} / kWh", body_style)],
            [Paragraph("Grid Share Contribution Ratio", body_style), Paragraph(str(f.grid_contribution_ratio if f else "N/A"), body_style)]
        ]
        
        t = Table(table_data, colWidths=[200, 200])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#EDF2F7')),
            ('ALIGN', (0,0), (-1,-1), 'LEFT'),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E0')),
            ('PADDING', (0,0), (-1,-1), 6),
        ]))
        elements.append(t)
        elements.append(Spacer(1, 15))
        
        elements.append(Paragraph("2. Financial Disclosures & Assumptions", section_style))
        elements.append(Paragraph(f"<b>Compounding Lifetime Degradation Rate Source</b>: {DEGRADATION_RATE_SOURCE}", body_style))
        elements.append(Spacer(1, 6))
        elements.append(Paragraph(f"<b>Retail Electricity Rate Price Source</b>: {ELECTRICITY_RATE_SOURCE}", disclosure_style))
        elements.append(Spacer(1, 6))
        elements.append(Paragraph(f"<b>Grid Contribution Share note</b>: {GRID_CONTRIBUTION_NOTE}", disclosure_style))

    else:
        raise ValueError("Invalid report type")
        
    doc.build(elements)
    buffer.seek(0)
    return buffer.getvalue()

def generate_excel_report(site_id: int, report_type: str, db: Session) -> bytes:
    data = fetch_report_data(site_id, db)
    site = data["site"]
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    
    title_font = Font(name="Calibri", size=16, bold=True, color="1A365D")
    section_font = Font(name="Calibri", size=12, bold=True, color="2B6CB0")
    header_font = Font(name="Calibri", size=11, bold=True, color="000000")
    bold_font = Font(name="Calibri", size=10, bold=True)
    body_font = Font(name="Calibri", size=10)
    disclosure_font = Font(name="Calibri", size=9, italic=True, color="4A5568")
    
    header_fill = PatternFill(start_color="EDF2F7", end_color="EDF2F7", fill_type="solid")
    disclosure_fill = PatternFill(start_color="F7FAFC", end_color="F7FAFC", fill_type="solid")
    
    # Title
    ws.cell(row=1, column=1, value=f"Site Report: {site.name}").font = title_font
    ws.cell(row=2, column=1, value=f"Coordinates: {data['lat']:.5f}, {data['lon']:.5f}").font = body_font
    ws.cell(row=3, column=1, value=f"Report: {report_type.replace('_', ' ').title()}").font = body_font
    
    curr_row = 5
    
    if report_type == "site_assessment":
        # Section 1
        ws.cell(row=curr_row, column=1, value="Deployment Suitability Breakdown").font = section_font
        curr_row += 1
        
        headers = ["Metric", "Score / Value"]
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=curr_row, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="left")
        curr_row += 1
        
        score = data["score"]
        score_val = score.overall_deployment_score if score else 0.0
        cat_val = score.suitability_category if score else "N/A"
        
        rows = [
            ("Overall Score", score_val),
            ("Suitability Category", cat_val),
            ("Solar Sub-score", score.solar_score if score else "N/A"),
            ("Wind Sub-score", score.wind_score if score else "N/A"),
            ("Geographic Score", score.geographic_score if score else "N/A"),
            ("Infrastructure Score", score.infrastructure_score if score else "N/A"),
            ("Environmental Score", score.environmental_score if score else "N/A"),
            ("Economic Score", score.economic_score if score else "N/A")
        ]
        for key, val in rows:
            ws.cell(row=curr_row, column=1, value=key).font = body_font
            ws.cell(row=curr_row, column=2, value=val).font = body_font
            curr_row += 1
            
        curr_row += 1
        # Section 2 Disclosures
        ws.cell(row=curr_row, column=1, value="Data Sources & Regulatory Disclosures").font = section_font
        curr_row += 1
        
        disclosures = [
            f"Formula Used: {SUITABILITY_FORMULA_USED} ({SUITABILITY_FORMULA_SOURCE})",
            f"Economic score proxy disclosure: {ECONOMIC_SCORE_NOTE}",
            f"Environmental score Copernicus blocker note: {ENVIRONMENTAL_SCORE_NOTE}",
            f"Infrastructure proximity proxy note: {INFRASTRUCTURE_SCORE_NOTE}"
        ]
        for note in disclosures:
            ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row+1, end_column=4)
            cell = ws.cell(row=curr_row, column=1, value=note)
            cell.font = disclosure_font
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.fill = disclosure_fill
            curr_row += 3
            
    elif report_type == "potential":
        # Section 1
        ws.cell(row=curr_row, column=1, value="Solar and Wind Generation Potential").font = section_font
        curr_row += 1
        
        headers = ["Resource Indicator", "Value"]
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=curr_row, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="left")
        curr_row += 1
        
        s_pred = data["solar_pred"]
        w_pred = data["wind_pred"]
        
        rows = [
            ("Annual Solar Incident Irradiance", f"{s_pred.annual_irradiance if s_pred else 'N/A'} kWh/m²"),
            ("Expected Annual Solar Output (1kW reference)", f"{s_pred.expected_energy_output if s_pred else 'N/A'} kWh/year"),
            ("Solar Capacity Factor", s_pred.capacity_factor if s_pred else "N/A"),
            ("Average Wind Speed", f"{w_pred.average_wind_speed if w_pred else 'N/A'} m/s"),
            ("Wind Power Density (WPD)", f"{w_pred.wind_power_density if w_pred else 'N/A'} W/m²"),
            ("Wind Turbulence Intensity (TI)", w_pred.turbulence_intensity if w_pred else "N/A"),
            ("Wind Capacity Factor (2MW reference)", w_pred.capacity_factor if w_pred else "N/A"),
            ("Expected Annual Wind Energy (2MW reference)", f"{w_pred.expected_annual_energy_production if w_pred else 'N/A'} kWh/year")
        ]
        for key, val in rows:
            ws.cell(row=curr_row, column=1, value=key).font = body_font
            ws.cell(row=curr_row, column=2, value=val).font = body_font
            curr_row += 1
            
        curr_row += 1
        # Section 2 Disclosures
        ws.cell(row=curr_row, column=1, value="Performance Disclosures & Physics Constants").font = section_font
        curr_row += 1
        
        disclosures = [
            f"Performance Formula: {FORECAST_FORMULA_USED} ({FORECAST_FORMULA_SOURCE})",
            f"Solar CF proxy note: {SOLAR_CAP_FACTOR_NOTE}",
            f"Wind CF proxy note: {WIND_CAP_FACTOR_NOTE}",
            f"Wind turbulence assumption: {WIND_TURB_FORMULA_USED} ({WIND_TURB_FORMULA_SOURCE})"
        ]
        for note in disclosures:
            ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row+1, end_column=4)
            cell = ws.cell(row=curr_row, column=1, value=note)
            cell.font = disclosure_font
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.fill = disclosure_fill
            curr_row += 3

    elif report_type == "feasibility":
        # Section 1
        ws.cell(row=curr_row, column=1, value="Economic Feasibility & Grid Contribution").font = section_font
        curr_row += 1
        
        headers = ["Economic Metric", "Value"]
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=curr_row, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="left")
        curr_row += 1
        
        f = data["forecast"]
        
        rows = [
            ("Solar Projected Year-1 Revenue", f"${f.solar_annual_revenue if f else 'N/A'}"),
            ("Wind Projected Year-1 Revenue", f"${f.wind_annual_revenue if f else 'N/A'}"),
            ("Combined Projected Year-1 Revenue", f"${f.combined_annual_revenue if f else 'N/A'}"),
            ("Electricity Price Utility Rate", f"${f.electricity_rate_usd_kwh if f else 'N/A'} / kWh"),
            ("Grid Share Contribution Ratio", f.grid_contribution_ratio if f else "N/A")
        ]
        for key, val in rows:
            ws.cell(row=curr_row, column=1, value=key).font = body_font
            ws.cell(row=curr_row, column=2, value=val).font = body_font
            curr_row += 1
            
        curr_row += 1
        # Section 2 Disclosures
        ws.cell(row=curr_row, column=1, value="Financial Disclosures & Assumptions").font = section_font
        curr_row += 1
        
        disclosures = [
            f"Compounding Lifetime Degradation Rate Source: {DEGRADATION_RATE_SOURCE}",
            f"Retail Electricity Rate Price Source: {ELECTRICITY_RATE_SOURCE}",
            f"Grid Contribution Share note: {GRID_CONTRIBUTION_NOTE}"
        ]
        for note in disclosures:
            ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row+1, end_column=4)
            cell = ws.cell(row=curr_row, column=1, value=note)
            cell.font = disclosure_font
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.fill = disclosure_fill
            curr_row += 3
            
    else:
        raise ValueError("Invalid report type")
        
    # Auto-fit columns
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
        
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
