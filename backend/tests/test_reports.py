import pytest
from io import BytesIO
from unittest.mock import patch
from sqlalchemy.orm import Session
from openpyxl import load_workbook

from app.models.user import User, Role
from app.models.project import Project
from app.models.site import Site
from app.models.site_scores import SiteScore
from app.models.environmental import SiteEnvironmentalData
from app.models.solar_prediction import SolarPrediction
from app.models.wind_prediction import WindPrediction
from app.models.energy_forecast import EnergyForecast

from app.services.report_generator import generate_pdf_report, generate_excel_report

# Disclosures
from app.services.suitability_engine import ENVIRONMENTAL_SCORE_NOTE
from app.api.predictions import WIND_CAP_FACTOR_NOTE
from app.services.forecasting_engine import GRID_CONTRIBUTION_NOTE

def setup_report_data(db_session: Session) -> Site:
    # 1. User & Project
    planner_role = db_session.query(Role).filter(Role.name == "Planner").first()
    owner = User(email="report_owner@example.com", hashed_password="pw", is_active=True)
    owner.roles.append(planner_role)
    db_session.add(owner)
    db_session.commit()
    
    project = Project(name="Report Project", owner_id=owner.id)
    db_session.add(project)
    db_session.commit()
    
    # 2. Site
    site = Site(project_id=project.id, name="Report Site", geom="SRID=4326;POINT(-115 35)")
    db_session.add(site)
    db_session.commit()
    
    # 3. Associated records
    score = SiteScore(
        site_id=site.id,
        overall_deployment_score=80.0,
        suitability_category="High Suitability",
        solar_score=85.0,
        wind_score=75.0,
        geographic_score=80.0,
        infrastructure_score=80.0,
        environmental_score=80.0,
        economic_score=80.0
    )
    env = SiteEnvironmentalData(
        site_id=site.id,
        average_solar_irradiance=5.0,
        average_wind_speed=6.0,
        nearest_road_km=1.0,
        nearest_substation_km=2.0,
        nearest_urban_area_km=5.0,
        nearest_protected_zone_km=3.0,
        nearest_water_body_km=4.0
    )
    solar_pred = SolarPrediction(
        site_id=site.id,
        annual_irradiance=1825.0,
        peak_sun_hours=5.0,
        expected_energy_output=1500.0,
        capacity_factor=0.17,
        performance_ratio=0.75
    )
    wind_pred = WindPrediction(
        site_id=site.id,
        average_wind_speed=6.0,
        wind_power_density=200.0,
        turbulence_intensity=0.12,
        capacity_factor=0.25,
        expected_annual_energy_production=4000000.0
    )
    forecast = EnergyForecast(
        site_id=site.id,
        solar_seasonal_kwh=[100] * 12,
        wind_seasonal_kwh=[300000] * 12,
        solar_longterm_kwh=[1500] * 20,
        wind_longterm_kwh=[4000000] * 20,
        solar_annual_revenue=187.5,
        wind_annual_revenue=500000.0,
        combined_annual_revenue=500187.5,
        electricity_rate_usd_kwh=0.125,
        grid_contribution_ratio=0.5
    )
    
    db_session.add_all([score, env, solar_pred, wind_pred, forecast])
    db_session.commit()
    return site

def test_pdf_export_contains_disclosure_text(db_session: Session):
    site = setup_report_data(db_session)
    
    # Spy/mock SimpleDocTemplate.build to examine the elements list
    with patch("app.services.report_generator.SimpleDocTemplate.build") as mock_build:
        generate_pdf_report(site.id, "site_assessment", db_session)
        elements = mock_build.call_args[0][0]
        
        # Check that environmental score Copernicus blocker note exists in elements
        found = False
        for el in elements:
            if hasattr(el, "text") and ENVIRONMENTAL_SCORE_NOTE in el.text:
                found = True
        assert found is True

    with patch("app.services.report_generator.SimpleDocTemplate.build") as mock_build:
        generate_pdf_report(site.id, "potential", db_session)
        elements = mock_build.call_args[0][0]
        
        # Check that wind capacity factor note exists in elements
        found = False
        for el in elements:
            if hasattr(el, "text") and WIND_CAP_FACTOR_NOTE in el.text:
                found = True
        assert found is True

    with patch("app.services.report_generator.SimpleDocTemplate.build") as mock_build:
        generate_pdf_report(site.id, "feasibility", db_session)
        elements = mock_build.call_args[0][0]
        
        # Check that grid contribution baseline note exists in elements
        found = False
        for el in elements:
            if hasattr(el, "text") and GRID_CONTRIBUTION_NOTE in el.text:
                found = True
        assert found is True

def test_excel_export_contains_disclosure_text(db_session: Session):
    site = setup_report_data(db_session)
    
    # 1. Site Assessment Report Excel check
    excel_bytes = generate_excel_report(site.id, "site_assessment", db_session)
    wb = load_workbook(BytesIO(excel_bytes))
    ws = wb.active
    
    found = False
    for row in ws.iter_rows(values_only=True):
        for val in row:
            if val and ENVIRONMENTAL_SCORE_NOTE in str(val):
                found = True
    assert found is True
    
    # 2. Potential Report Excel check
    excel_bytes = generate_excel_report(site.id, "potential", db_session)
    wb = load_workbook(BytesIO(excel_bytes))
    ws = wb.active
    
    found = False
    for row in ws.iter_rows(values_only=True):
        for val in row:
            if val and WIND_CAP_FACTOR_NOTE in str(val):
                found = True
    assert found is True

    # 3. Feasibility Report Excel check
    excel_bytes = generate_excel_report(site.id, "feasibility", db_session)
    wb = load_workbook(BytesIO(excel_bytes))
    ws = wb.active
    
    found = False
    for row in ws.iter_rows(values_only=True):
        for val in row:
            if val and GRID_CONTRIBUTION_NOTE in str(val):
                found = True
    assert found is True
